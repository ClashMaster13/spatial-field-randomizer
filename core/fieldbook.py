import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

def get_color_palette(n):
    palette = [
        ("BBDEFB", "000000"), ("FFE0B2", "000000"), ("C8E6C9", "000000"), ("FFCDD2", "000000"),
        ("D1C4E9", "000000"), ("FFF9C4", "000000"), ("B2DFDB", "000000"), ("F8BBD0", "000000"),
        ("DCEDC8", "000000"), ("B3E5FC", "000000"), ("FFE082", "000000"), ("D7CCC8", "000000"),
        ("CFD8DC", "000000"), ("E1BEE7", "000000"), ("FFCCBC", "000000"), ("C5CAE9", "000000"),
        ("B2EBF2", "000000"), ("E6EE9C", "000000"), ("F0F4C3", "000000"), ("E0F2F1", "000000"),
        ("EDE7F6", "000000"), ("FFF3E0", "000000"), ("E8F5E9", "000000"), ("FCE4EC", "000000"),
        ("E3F2FD", "000000"), ("FFF8E1", "000000"), ("F3E5F5", "000000"), ("FBE9E7", "000000"),
        ("EFEBE9", "000000"), ("ECEFF1", "000000")
    ]
    if n <= len(palette): return palette[:n]
    extended = list(palette)
    for _ in range(len(palette), n): extended.append(("F5F5F5", "000000"))
    return extended

def generate_excel_bytes(trial_name, genotypes, final_grids, reps, rows, cols, start_plot, start_row, start_col):
    # 1. Map plots
    plot_map = {g: {} for g in genotypes}
    curr_p = start_plot

    for i in range(reps):
        r_num = i + 1
        grid = final_grids[i]
        for r in range(rows):
            global_r = start_row + (i * rows) + r
            col_idx = list(range(cols)) if (i * rows + r) % 2 == 0 else list(range(cols - 1, -1, -1))
            for c in col_idx:
                g = grid[r][c]
                plot_map[g][r_num] = {"plot": curr_p, "row": global_r, "col": start_col + c}
                curr_p += 1

    # 2. Build workbook
    colors = get_color_palette(len(genotypes))
    color_dict = {g: colors[idx] for idx, g in enumerate(genotypes)}

    wb = openpyxl.Workbook()
    tnr_norm = Font(name="Times New Roman")
    tnr_bold = Font(name="Times New Roman", bold=True)
    tnr_head = Font(name="Times New Roman", bold=True, color="FFFFFF")
    head_fill = PatternFill(start_color="337AB7", end_color="337AB7", fill_type="solid")
    align_c = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

    # Sheet 1: Fieldbook
    ws1 = wb.active
    ws1.title = f"{trial_name} Fieldbook"[:31]
    r_idx = 1

    for rep in range(1, reps + 1):
        ws1.cell(row=r_idx, column=1, value=f"REP {rep} ORDER").font = tnr_bold
        r_idx += 1
        others = [x for x in range(1, reps + 1) if x != rep]
        headers = ["S. No.", "Genotype", f"R{rep} Row", f"R{rep} Col", f"R{rep} Plot"] + [f"R{x} Plot" for x in others]

        for c_idx, h in enumerate(headers, 1):
            c = ws1.cell(row=r_idx, column=c_idx, value=h)
            c.font, c.fill, c.alignment, c.border = tnr_head, head_fill, align_c, border_thin
        r_idx += 1

        sorted_lines = sorted(genotypes, key=lambda x: plot_map[x][rep]["plot"])
        for i, g in enumerate(sorted_lines, 1):
            data = plot_map[g][rep]
            row_data = [i, g, data["row"], data["col"], data["plot"]] + [plot_map[g][x]["plot"] for x in others]
            for c_idx, val in enumerate(row_data, 1):
                c = ws1.cell(row=r_idx, column=c_idx, value=val)
                c.font, c.alignment, c.border = tnr_norm, align_c, border_thin
            r_idx += 1
        r_idx += 2

    for c in ws1.columns: ws1.column_dimensions[c[0].column_letter].width = 12

    # Sheet 2: Spatial Map
    ws2 = wb.create_sheet(title=f"{trial_name} Grid Map"[:31])
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols + 2)
    ws2.cell(row=1, column=1, value=f"{trial_name} Spatial Map").font = Font(name="Times New Roman", bold=True, size=14)
    ws2.cell(row=1, column=1).alignment = align_c

    thick = Side(style="medium")
    thin = Side(style="thin")
    map_r = 3

    for rep in range(reps - 1, -1, -1):
        start_map_r = map_r
        grid = final_grids[rep]
        for r in range(rows - 1, -1, -1):
            global_r = start_row + (rep * rows) + r
            ws2.cell(row=map_r, column=2, value=f"Row {global_r}").font = tnr_bold
            ws2.cell(row=map_r, column=2).alignment = align_c
            for c in range(cols):
                g = grid[r][c]
                p_no = plot_map[g][rep + 1]["plot"]
                cell = ws2.cell(row=map_r, column=c + 3, value=f"Plot {p_no}\n{g}")
                bg, fg = color_dict[g]
                cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
                cell.font = Font(name="Times New Roman", bold=True, color=fg)
                cell.alignment = align_c
                ws2.column_dimensions[cell.column_letter].width = 15

            ws2.row_dimensions[map_r].height = 45
            map_r += 1

        ws2.merge_cells(start_row=start_map_r, start_column=1, end_row=map_r - 1, end_column=1)
        r_cell = ws2.cell(row=start_map_r, column=1, value=f"R {rep + 1}")
        r_cell.font, r_cell.alignment = tnr_bold, align_c

    ws2.row_dimensions[map_r].height = 20
    for c in range(cols):
        c_cell = ws2.cell(row=map_r, column=c + 3, value=f"Col {start_col + c}")
        c_cell.font, c_cell.alignment = tnr_bold, align_c

    for r in range(3, map_r + 1):
        for c in range(1, cols + 3):
            cell = ws2.cell(row=r, column=c)
            if r == map_r and c < 3: continue
            
            t, b, l, re = thin, thin, thin, thin
            if r == 3: t = thick
            if r == map_r - 1: b = thick
            if r == map_r: t, b = thick, thick
            
            idx = r - 3
            if idx % rows == 0 and r != 3 and r < map_r: t = thick
            if (idx + 1) % rows == 0 and r < map_r: b = thick
            
            if c in [1, 2]: l, re = thick, thick
            elif c == 3: l = thick
            elif c == cols + 2: re = thick
                
            cell.border = Border(top=t, bottom=b, left=l, right=re)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()